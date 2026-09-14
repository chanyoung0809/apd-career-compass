# -*- coding: utf-8 -*-
"""APD1 1:1 진로상담 기록 양식 자동 생성

용도
- APD1(AI 프로덕트 개발 트랙 1기) 1차 진로상담(1인당 30분)을 위해 통합 1파일 2시트 양식을 생성합니다.
  - 시트 1 '상담 로그'  : 11명 한 줄씩. 사전 추정 진로·핵심 목표는 자동 채움, 상담 기록 칸은 빈칸.
  - 시트 2 '레이서 프로필': 지원동기·수료후목표 원문. 상담 직전 읽기용.
- 데이터 출처: [APD1]레이서_정보.xlsx (OneDrive KDT_공유폴더 12. APD 트랙\\APD 1기)

주의
- 이 스크립트를 다시 실행하면 출력 파일이 덮어쓰기됩니다.
- 상담 메모가 입력된 후에는 재실행을 피하세요. 필요 시 기존 파일을 백업 후 실행.
"""
import glob
import os
import sys
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

# KDT 공유폴더(OneDrive) 안에서 레이서 정보 파일을 자동 탐색
_SHARE = glob.glob(os.path.join(os.path.expanduser("~"), "OneDrive - 엘리스", "*KDT_공유폴더*"))
if not _SHARE:
    raise SystemExit("OneDrive KDT_공유폴더를 찾지 못했습니다. 폴더명을 확인하세요.")
RACER_XLSX = os.path.join(
    _SHARE[0], "02. KDT 운영", "12. APD 트랙", "APD 1기", "[APD1]레이서_정보.xlsx"
)
OUT_XLSX = r"C:\KDT AX\APD1_OneOnOne\[APD1] 1on1 상담 기록.xlsx"

# 레이서 정보 시트의 수료후목표를 상담용으로 요약한 것 (상담 로그 B영역 자동 채움용)
CORE_GOAL = {
    "김가을": "배운 내용을 프로젝트에 적용, 포트폴리오 완성 후 새 직무로 전환",
    "김서연": "AI 활용 서비스기획·PM으로 커리어 확장 (메디컬 도메인 특화)",
    "김호림": "AI 서비스 개발 관련 직무 취업 (가장 우선적인 목표)",
    "박서현": "지역 상생 플랫폼 '가치잇다'를 MVP로 구현 후 사용자 검증",
    "서지훈": "직접 설계해 온 ERP 시스템을 운영 가능한 AI 서비스로 완성",
    "이현진": "머리 속 아이디어를 직접 실행해 저만의 프로덕트 개발",
    "전세황": "프론트(React/Vue)·백엔드(FastAPI)·DB 실무 역량으로 웹앱 개발 취업",
    "정우현": "AI 서비스 개발 실무 역량·프로젝트 경험을 갖춰 개발 직무 취업",
    "최은비": "AI 개발 기술×비즈니스 설계로 1인 창업 MVP를 직접 개발·검증",
    "홍태휘": "취업 최종합격 (구체 계획 미정 → 상담에서 목표 정리 필요)",
    "황지혜": "기획에서 멈추지 않고 직접 설계·구현해 출시까지 해보는 것",
}

# 지원서·수강목적 기준 사전 추정 진로 (상담에서 확정 칸을 따로 둠)
PRE_ESTIMATE = {
    "김가을": "미확정 (이직 목표, 직무 미정)",
    "김서연": "PM (서비스기획 전환 희망)",
    "김호림": "개발자 (AI 서비스 개발 취업)",
    "박서현": "창업 ('가치잇다' MVP)",
    "서지훈": "창업 (AI 서비스 완성)",
    "이현진": "개발자 (풀스택, 창업 병행)",
    "전세황": "개발자 (웹앱 개발 취업)",
    "정우현": "개발자 (AI 서비스 개발 취업)",
    "최은비": "창업 (1인 MVP, PM 병행)",
    "홍태휘": "미확정 (취업 목표)",
    "황지혜": "미확정·복수 (개발·PM 모두 흥미)",
}

HEAD_FILL = PatternFill("solid", fgColor="2BB8EA")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="C9D8E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
WRAP_CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)


def load_racers():
    import pandas as pd

    df = pd.read_excel(RACER_XLSX, header=1)
    df = df[df["이름"].notna()]
    rows = []
    for _, r in df.iterrows():
        name = str(r["이름"]).strip()
        rows.append({
            "이름": name,
            "전공": fmt(r["학과/전공"]),
            "재직 상태": fmt(r["재직 상태"]),
            "현재 직무": fmt(r["현재 직무"]),
            "희망 직무": fmt(r["희망 직무"]),
            "수강목적": fmt(r["수강목적"]),
            "지원동기": fmt(r["지원동기"]),
            "수료후목표": fmt(r["수료후목표"]),
        })
    return rows


def fmt(v, default="—"):
    s = "" if v is None else str(v).strip()
    if not s or s.lower() == "nan":
        return default
    return s.replace("\r\n", "\n")


def style_sheet(ws, widths, height_first_col=2):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
            cell.border = BORDER
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 30


def build_log_sheet(wb, racers):
    ws = wb.create_sheet("상담 로그")
    headers = [
        "연번", "이름", "희망 진로\n(사전 추정)", "희망 진로\n(상담 확정)",
        "현재 상황", "전공", "수강목적", "핵심 목표 (수료후)",
        "상담 일시", "상담 내용", "고민 키워드", "액션 아이템", "후속 조치",
    ]
    ws.append(headers)
    for i, r in enumerate(racers, start=1):
        status = r["재직 상태"] if r["재직 상태"] != "—" else "취업 준비생"
        if r["현재 직무"] != "—":
            status += f" · {r['현재 직무']}"
        ws.append([
            i, r["이름"], PRE_ESTIMATE.get(r["이름"], "—"), None,
            status, r["전공"], r["수강목적"], CORE_GOAL.get(r["이름"], "—"),
            None, None, None, None, None,
        ])
    style_sheet(ws, [5, 9, 20, 14, 16, 12, 12, 30, 12, 36, 16, 30, 16])
    for idx in range(2, len(racers) + 2):
        ws.row_dimensions[idx].height = 54
    note = ws.cell(row=len(racers) + 3, column=2,
                   value="※ '상담 일시' 이후 칸은 상담 중 직접 기입. '희망 진로(상담 확정)'는 개발자/PM/창업/미확정 중 하나로.")
    note.font = Font(color="73838E", italic=True, size=10)
    note.alignment = Alignment(vertical="center")


def build_profile_sheet(wb, racers):
    ws = wb.create_sheet("레이서 프로필")
    headers = [
        "연번", "이름", "전공", "재직 상태", "현재 직무",
        "희망 직무", "수강목적", "지원동기 (원문)", "수료후목표 (원문)",
    ]
    ws.append(headers)
    for i, r in enumerate(racers, start=1):
        ws.append([
            i, r["이름"], r["전공"], r["재직 상태"], r["현재 직무"],
            r["희망 직무"], r["수강목적"], r["지원동기"], r["수료후목표"],
        ])
    style_sheet(ws, [5, 9, 12, 12, 12, 16, 12, 60, 60])
    for idx in range(2, len(racers) + 2):
        ws.row_dimensions[idx].height = 120


def main():
    racers = load_racers()
    print(f"레이서 {len(racers)}명 로드: " + ", ".join(r["이름"] for r in racers))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_log_sheet(wb, racers)
    build_profile_sheet(wb, racers)
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"생성 완료: {OUT_XLSX}")


if __name__ == "__main__":
    main()
