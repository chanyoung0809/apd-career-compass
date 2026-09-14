# -*- coding: utf-8 -*-
"""APD1 1:1 상담 입력 로컬 서버

- 상담입력.html 을 http://localhost:8765 로 서빙하고,
  브라우저에서 입력한 상담 기록을 [APD1] 1on1 상담 기록.xlsx 에 바로 저장합니다.
- GET  /api/data : 개인별 시트의 라벨·값·읽기전용 정보를 JSON으로 반환
- POST /api/save : {"updates": {시트명: {라벨: 값}}} 을 받아 xlsx에 반영
  (개인별 시트와 '상담 로그' 시트를 함께 갱신해 두 시트가 항상 일치)
- 실행: python tools/상담_서버.py  → 브라우저 자동 오픈
- 주의: xlsx가 Excel에서 열려 있으면 저장이 막힙니다. 저장 전 엑셀 창을 닫아 주세요.
"""
import io
import json
import os
import sys
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "[APD1] 1on1 상담 기록.xlsx")
HTML = os.path.join(BASE, "상담입력.html")
PORT = 8765

# 개인별 시트에서 브라우저로 편집할 라벨 (A열 라벨 ↔ B열 값)
EDITABLE = [
    "상담 일시", "시작 시간", "소요 시간", "희망 진로 (상담 확정)", "한 줄 요약",
    "질문 1 (끌리는 진로)", "답 1", "질문 2 (가장 큰 걱정)", "답 2",
    "질문 3 (꼭 물어볼 질문)", "답 3",
    "상담 내용", "고민 키워드", "액션 아이템", "후속 조치",
]
# '상담 로그' 시트에도 함께 반영할 필드 (키=개인별 라벨, 값=로그 헤더·개행 정규화 기준)
LOG_SYNC = {
    "상담 일시": "상담 일시",
    "시작 시간": "시작 시간",
    "소요 시간": "소요 시간",
    "희망 진로 (상담 확정)": "희망 진로 (상담 확정)",
    "상담 내용": "상담 내용",
    "고민 키워드": "고민 키워드",
    "액션 아이템": "액션 아이템",
    "후속 조치": "후속 조치",
}
READONLY = [
    "1차 프로젝트", "현재 학습 구간", "커리큘럼 순서 참고",
    "전공", "재직 상태·현재 직무", "희망 직무", "수강목적",
    "핵심 목표 (수료후)", "지원동기 (원문)", "수료후목표 (원문)",
]

SAVE_LOCK = threading.Lock()


def label_rows(ws):
    """A열 라벨 → 행 번호"""
    return {
        str(ws.cell(row=i, column=1).value).strip(): i
        for i in range(1, ws.max_row + 1)
        if ws.cell(row=i, column=1).value is not None
    }


def load_data():
    wb = openpyxl.load_workbook(XLSX)
    racers = []
    for ws in wb.worksheets:
        if not (ws.title[:2].isdigit() and "_" in ws.title):
            continue
        rows = label_rows(ws)
        if "이름" not in rows:
            continue
        name = str(ws.cell(row=rows["이름"], column=2).value or "").strip()
        type_page = None
        if "맞춤 유형 페이지" in rows:
            c = ws.cell(row=rows["맞춤 유형 페이지"], column=2)
            type_page = {"text": c.value or "",
                         "url": c.hyperlink.target if c.hyperlink else None}
        guide = []
        for i in range(1, ws.max_row + 1):
            a = ws.cell(row=i, column=1).value
            if isinstance(a, str) and a.strip().startswith("☐"):
                guide.append(ws.cell(row=i, column=2).value or "")
        tab = ws.sheet_properties.tabColor
        racers.append({
            "sheet": ws.title,
            "name": name,
            "no": int(ws.title[:2]),
            "pre": ws.cell(row=rows["희망 진로 (사전 추정)"], column=2).value or "",
            "tabColor": ("#" + tab.rgb[-6:]) if tab and tab.rgb else None,
            "typePage": type_page,
            "fields": {lab: (ws.cell(row=rows[lab], column=2).value or "")
                       for lab in EDITABLE if lab in rows},
            "info": {lab: (ws.cell(row=rows[lab], column=2).value or "")
                     for lab in READONLY if lab in rows},
            "guide": guide,
        })
    racers.sort(key=lambda r: r["no"])
    return {"racers": racers}


def apply_save(payload):
    updates = payload["updates"]
    wb = openpyxl.load_workbook(XLSX)
    log = wb["상담 로그"]

    # 로그 헤더(개행 정규화) → 열 번호, 이름 → 행 번호
    headers = {}
    for c in range(1, log.max_column + 1):
        v = log.cell(row=1, column=c).value
        if v is not None:
            headers[str(v).replace("\n", " ").strip()] = c
    name_row = {}
    for r in range(2, log.max_row + 1):
        v = log.cell(row=r, column=2).value
        if v is not None:
            name_row[str(v).strip()] = r

    changed = []
    for sheet_name, field_updates in updates.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = label_rows(ws)
        if "이름" not in rows:
            continue
        name = str(ws.cell(row=rows["이름"], column=2).value or "").strip()
        for label, value in field_updates.items():
            if label not in rows:
                continue
            ws.cell(row=rows[label], column=2).value = value or None
            log_label = LOG_SYNC.get(label)
            if log_label and log_label in headers and name in name_row:
                log.cell(row=name_row[name],
                         column=headers[log_label]).value = value or None
        changed.append(sheet_name)

    # 임시 파일 저장 후 원자적 교체 (쓰기 도중 손상 방지)
    tmp = XLSX + ".tmp"
    wb.save(tmp)
    os.replace(tmp, XLSX)
    return changed


class Handler(BaseHTTPRequestHandler):
    def _send(self, code, body, ctype="application/json; charset=utf-8"):
        data = body.encode("utf-8") if isinstance(body, str) else body
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _json(self, code, obj):
        self._send(code, json.dumps(obj, ensure_ascii=False))

    def do_GET(self):
        if self.path in ("/", "/index.html", "/상담입력.html"):
            try:
                with open(HTML, "rb") as f:
                    self._send(200, f.read(), "text/html; charset=utf-8")
            except OSError:
                self._json(500, {"error": "상담입력.html 을 찾지 못했습니다."})
        elif self.path == "/api/data":
            try:
                self._json(200, load_data())
            except PermissionError:
                self._json(503, {"error": "엑셀 파일이 다른 프로그램에 잠겨 있습니다. "
                                          "Excel 창을 닫고 새로고침해 주세요."})
            except Exception as e:
                self._json(500, {"error": f"{type(e).__name__}: {e}"})
        else:
            self._json(404, {"error": "not found"})

    def do_POST(self):
        if self.path != "/api/save":
            return self._json(404, {"error": "not found"})
        try:
            length = int(self.headers.get("Content-Length", 0))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            with SAVE_LOCK:
                changed = apply_save(payload)
            self._json(200, {"ok": True, "changed": changed})
        except PermissionError:
            self._json(423, {"ok": False, "error": "엑셀 파일이 열려 있어 저장하지 못했어요. "
                             "Excel 창을 닫은 뒤 다시 '저장'을 눌러 주세요."})
        except Exception as e:
            self._json(500, {"ok": False, "error": f"{type(e).__name__}: {e}"})

    def log_message(self, *args):  # 콘솔 소음 줄이기
        pass


def main():
    if not os.path.exists(XLSX):
        raise SystemExit("[APD1] 1on1 상담 기록.xlsx 을 찾지 못했습니다. "
                         "python tools/make_상담기록.py 를 먼저 실행하세요.")
    try:
        server = ThreadingHTTPServer(("127.0.0.1", PORT), Handler)
    except OSError:
        raise SystemExit(f"포트 {PORT} 가 이미 사용 중입니다. "
                         "이미 실행 중인 상담 서버가 있는지 확인하세요.")
    url = f"http://localhost:{PORT}"
    print(f"상담 입력 페이지: {url}   (종료: Ctrl+C)")
    print(f"대상 엑셀 파일: {XLSX}")
    threading.Timer(0.5, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n서버 종료")


if __name__ == "__main__":
    main()
